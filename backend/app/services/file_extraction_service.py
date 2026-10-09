import csv
import os

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
import pytesseract
from PIL import Image

from app.services.pdf_service import extract_text_from_pdf
from app.services.ocr_service import extract_text_from_scanned_pdf


# ============================================================
# TXT EXTRACTION
# ============================================================

def extract_text_from_txt(file_path: str):
    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"TXT file not found: {file_path}"
        )

    try:
        with open(
            file_path,
            "r",
            encoding="utf-8",
        ) as file:
            text = file.read()

    except UnicodeDecodeError:
        with open(
            file_path,
            "r",
            encoding="latin-1",
        ) as file:
            text = file.read()

    text = text.strip()

    return {
        "text": text,
        "pages": (
            [{"page_number": 1, "text": text}]
            if text
            else []
        ),
        "page_count": 1 if text else 0,
        "pages_with_text": 1 if text else 0,
        "is_scanned": False,
    }


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_text_from_docx(file_path: str):
    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"DOCX file not found: {file_path}"
        )

    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    text = "\n\n".join(paragraphs).strip()

    return {
        "text": text,
        "pages": (
            [{"page_number": 1, "text": text}]
            if text
            else []
        ),
        "page_count": 1 if text else 0,
        "pages_with_text": 1 if text else 0,
        "is_scanned": False,
    }


# ============================================================
# CSV EXTRACTION
# ============================================================

def extract_text_from_csv(file_path: str):
    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"CSV file not found: {file_path}"
        )

    rows = []
    last_error = None

    for encoding in [
        "utf-8-sig",
        "utf-8",
        "latin-1",
    ]:
        try:
            with open(
                file_path,
                "r",
                encoding=encoding,
                newline="",
            ) as file:
                reader = csv.reader(file)

                for row in reader:
                    cleaned_row = [
                        str(cell).strip()
                        for cell in row
                    ]

                    if any(cleaned_row):
                        rows.append(
                            " | ".join(cleaned_row)
                        )

            last_error = None
            break

        except UnicodeDecodeError as error:
            last_error = error

    if last_error:
        raise ValueError(
            "Unable to decode CSV file"
        ) from last_error

    text = "\n".join(rows).strip()

    return {
        "text": text,
        "pages": (
            [{"page_number": 1, "text": text}]
            if text
            else []
        ),
        "page_count": 1 if text else 0,
        "pages_with_text": 1 if text else 0,
        "is_scanned": False,
    }


# ============================================================
# XLSX EXTRACTION
# ============================================================

def extract_text_from_xlsx(file_path: str):
    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"XLSX file not found: {file_path}"
        )

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    pages = []
    combined_text_parts = []

    try:
        for sheet_number, worksheet in enumerate(
            workbook.worksheets,
            start=1,
        ):
            sheet_rows = []

            for row in worksheet.iter_rows(
                values_only=True
            ):
                values = []

                for value in row:
                    if value is None:
                        values.append("")
                    else:
                        values.append(
                            str(value).strip()
                        )

                if any(values):
                    sheet_rows.append(
                        " | ".join(values)
                    )

            if not sheet_rows:
                continue

            sheet_text = "\n".join(
                sheet_rows
            ).strip()

            if not sheet_text:
                continue

            page_text = (
                f"Sheet: {worksheet.title}\n"
                f"{sheet_text}"
            )

            pages.append(
                {
                    "page_number": sheet_number,
                    "text": page_text,
                }
            )

            combined_text_parts.append(
                page_text
            )

        sheet_count = len(workbook.worksheets)

    finally:
        workbook.close()

    combined_text = "\n\n".join(
        combined_text_parts
    ).strip()

    return {
        "text": combined_text,
        "pages": pages,
        "page_count": sheet_count,
        "pages_with_text": len(pages),
        "is_scanned": False,
    }


# ============================================================
# PPTX EXTRACTION
# ============================================================

def extract_text_from_pptx(file_path: str):
    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"PPTX file not found: {file_path}"
        )

    presentation = Presentation(file_path)

    pages = []
    combined_text_parts = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_parts = []

        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue

            text = shape.text.strip()

            if text:
                slide_parts.append(text)

        slide_text = "\n".join(
            slide_parts
        ).strip()

        if not slide_text:
            continue

        page_text = (
            f"Slide {slide_number}\n"
            f"{slide_text}"
        )

        pages.append(
            {
                "page_number": slide_number,
                "text": page_text,
            }
        )

        combined_text_parts.append(
            page_text
        )

    combined_text = "\n\n".join(
        combined_text_parts
    ).strip()

    return {
        "text": combined_text,
        "pages": pages,
        "page_count": len(
            presentation.slides
        ),
        "pages_with_text": len(pages),
        "is_scanned": False,
    }


# ============================================================
# IMAGE OCR
# ============================================================

def extract_text_from_image(file_path: str):
    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"Image file not found: {file_path}"
        )

    try:
        image = Image.open(file_path)

        text = pytesseract.image_to_string(
            image,
            lang="eng",
        ).strip()

    except Exception as error:
        print("Image OCR failed:")
        print(error)

        raise RuntimeError(
            "Unable to extract text from image"
        ) from error

    return {
        "text": text,
        "pages": (
            [{"page_number": 1, "text": text}]
            if text
            else []
        ),
        "page_count": 1,
        "pages_with_text": 1 if text else 0,
        "is_scanned": True,
    }


# ============================================================
# UNIFIED FILE EXTRACTION
# ============================================================

def extract_text_from_file(
    file_path: str,
    file_extension: str,
):
    extension = file_extension.lower().strip()

    if not extension.startswith("."):
        extension = "." + extension

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if extension == ".pdf":
        result = extract_text_from_pdf(
            file_path
        )

        if result["is_scanned"]:
            print(
                "PDF appears to be scanned. "
                "Starting OCR..."
            )

            ocr_result = (
                extract_text_from_scanned_pdf(
                    file_path
                )
            )

            return {
                "text": ocr_result["text"],
                "pages": ocr_result["pages"],
                "page_count": ocr_result[
                    "page_count"
                ],
                "pages_with_text": len(
                    ocr_result["pages"]
                ),
                "is_scanned": True,
            }

        return result

    # --------------------------------------------------------
    # TXT
    # --------------------------------------------------------

    if extension == ".txt":
        return extract_text_from_txt(
            file_path
        )

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    if extension == ".docx":
        return extract_text_from_docx(
            file_path
        )

    # --------------------------------------------------------
    # CSV
    # --------------------------------------------------------

    if extension == ".csv":
        return extract_text_from_csv(
            file_path
        )

    # --------------------------------------------------------
    # XLSX
    # --------------------------------------------------------

    if extension == ".xlsx":
        return extract_text_from_xlsx(
            file_path
        )

    # --------------------------------------------------------
    # PPTX
    # --------------------------------------------------------

    if extension == ".pptx":
        return extract_text_from_pptx(
            file_path
        )

    # --------------------------------------------------------
    # IMAGES
    # --------------------------------------------------------

    if extension in {
        ".png",
        ".jpg",
        ".jpeg",
    }:
        return extract_text_from_image(
            file_path
        )

    # --------------------------------------------------------
    # UNSUPPORTED
    # --------------------------------------------------------

    raise ValueError(
        f"Unsupported file type: {extension}"
    )